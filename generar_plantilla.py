"""
Genera la plantilla de datos ``datos/plantilla_datos.xlsx`` (y su versión CSV
en ``datos/ejemplo_csv/``) con un ejemplo de 15 trabajos.

Úsela como punto de partida: abra el archivo, reemplace los valores de
ejemplo por los suyos y guárdelo con otro nombre (p. ej. ``datos/mis_datos.xlsx``).

    python generar_plantilla.py
"""

import os

import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from hfs_ga.datos import NOMBRES_ETAPAS

CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "datos")


def construir_ejemplo(semilla: int = 7) -> dict[str, pd.DataFrame]:
    rng = np.random.default_rng(semilla)
    n = 15
    trabajos = [f"J{i + 1:02d}" for i in range(n)]

    cantidad = rng.integers(10, 61, n) * 10            # pares por lote
    fecha = np.round(rng.uniform(30, 100, n), 1)         # horas
    df_trab = pd.DataFrame({
        "trabajo": trabajos,
        "cantidad": cantidad,
        "fecha_entrega_h": fecha,
        "max_sublotes": 5,
    })

    # horas por unidad (a velocidad 1) en cada etapa
    df_t = pd.DataFrame({
        "trabajo": trabajos,
        "etapa_1": np.round(rng.uniform(0.010, 0.020, n), 4),   # troquelado
        "etapa_2": np.round(rng.uniform(0.150, 0.300, n), 4),   # ensamble
        "etapa_3": np.round(rng.uniform(0.060, 0.120, n), 4),   # costura
        "etapa_4": np.round(rng.uniform(0.050, 0.100, n), 4),   # volteado
    })

    filas = []
    config = [(1, "T", 4, False), (2, "E", 20, True), (3, "C", 9, False), (4, "V", 8, False)]
    for etapa, pref, m, uniforme in config:
        for i in range(m):
            filas.append({
                "etapa": etapa,
                "maquina": f"{pref}{i + 1:02d}",
                "velocidad": round(float(rng.uniform(0.8, 1.3)), 2) if uniforme else 1.0,
                "disponibilidad_h": round(float(rng.choice([0, 0, 0, 1, 2, 4, 6])), 1),
            })
    df_m = pd.DataFrame(filas)

    df_tr = pd.DataFrame({
        "de_etapa": [1, 2, 3],
        "a_etapa": [2, 3, 4],
        "tiempo_h": [0.5, 0.3, 0.4],
        "trabajo": ["", "", ""],
    })
    return {"Trabajos": df_trab, "Tiempos": df_t, "Maquinas": df_m, "Transporte": df_tr}


INSTRUCCIONES = [
    ("Hoja", "Columna", "Descripción"),
    ("Trabajos", "trabajo", "Identificador único del trabajo / lote (texto)."),
    ("Trabajos", "cantidad", "Número de unidades (p. ej. pares) del lote. Entero > 0."),
    ("Trabajos", "fecha_entrega_h", "Fecha de entrega d_j en HORAS desde el instante 0."),
    ("Trabajos", "max_sublotes", "(Opcional) Máximo de sublotes en que se puede dividir el lote. Por defecto 3."),
    ("Tiempos", "trabajo", "Debe coincidir con la hoja Trabajos."),
    ("Tiempos", "etapa_1 … etapa_4",
     "Tiempo de proceso por UNIDAD en HORAS (min/60). Etapas: " + ", ".join(NOMBRES_ETAPAS)
     + ". En la etapa 2 (uniforme) es el tiempo estándar a velocidad 1."),
    ("Maquinas", "etapa", "Número de etapa (1 a 4)."),
    ("Maquinas", "maquina", "Nombre de la máquina."),
    ("Maquinas", "velocidad", "Velocidad relativa. Tiempo real = tiempo unitario × cantidad / velocidad. "
                              "Usar 1 en etapas de máquinas idénticas (1, 3 y 4)."),
    ("Maquinas", "disponibilidad_h", "Instante (h) a partir del cual la máquina está disponible (am)."),
    ("Transporte", "de_etapa / a_etapa", "Par de etapas consecutivas (1→2, 2→3, 3→4)."),
    ("Transporte", "tiempo_h", "Tiempo de transporte de cada sublote entre esas etapas (h)."),
    ("Transporte", "trabajo", "(Opcional) Vacío = aplica a todos. Con un trabajo = valor específico para él."),
]


def guardar(tablas: dict[str, pd.DataFrame], ruta_xlsx: str, carpeta_csv: str | None = None) -> None:
    os.makedirs(os.path.dirname(ruta_xlsx), exist_ok=True)
    with pd.ExcelWriter(ruta_xlsx, engine="openpyxl") as w:
        pd.DataFrame(INSTRUCCIONES[1:], columns=INSTRUCCIONES[0]).to_excel(
            w, sheet_name="Instrucciones", index=False)
        for hoja, df in tablas.items():
            df.to_excel(w, sheet_name=hoja, index=False)
        for ws in w.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="305496")
            for i, col in enumerate(ws.columns, start=1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 90)
    if carpeta_csv:
        os.makedirs(carpeta_csv, exist_ok=True)
        for hoja, df in tablas.items():
            df.to_csv(os.path.join(carpeta_csv, hoja.lower() + ".csv"), index=False)


if __name__ == "__main__":
    tablas = construir_ejemplo()
    ruta = os.path.join(CARPETA, "plantilla_datos.xlsx")
    guardar(tablas, ruta, os.path.join(CARPETA, "ejemplo_csv"))
    print(f"Plantilla creada en: {ruta}")
    print(f"Versión CSV en:      {os.path.join(CARPETA, 'ejemplo_csv')}")
