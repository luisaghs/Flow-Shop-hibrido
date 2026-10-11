# =============================================================================
#  PROGRAMACIÓN DE LA PRODUCCIÓN EN UN FLOW SHOP HÍBRIDO DE 4 ETAPAS
#  ALGORITMO GENÉTICO + REGLAS DE DESPACHO
#
#      FH4, ((P4), (Q20), (P9), (P8)) | ls | ΣTj
#
#  Etapa 1: Troquelado  -> 4 máquinas paralelas idénticas   (P)
#  Etapa 2: Ensamble    -> 20 máquinas paralelas uniformes  (Q)
#  Etapa 3: Costura     -> 9 máquinas paralelas idénticas   (P)
#  Etapa 4: Volteado    -> 8 máquinas paralelas idénticas   (P)
#
#  ls   : lot streaming con sublotes consistentes (mismo tamaño en todas las etapas)
#  ΣTj  : se minimiza la tardanza total, Tj = max(0, Cj - dj)
#
#  Solo se consideran los tiempos de proceso de las 4 etapas; los tiempos de
#  transporte, alistamiento y demás se desprecian.
#
#  Los datos se leen de una BASE DE DATOS en Excel (ver hoja LEEME del archivo
#  base_datos_flow_shop.xlsx). El número de etapas y de máquinas se toma de la
#  hoja "Etapas", así que el mismo código sirve para instancias pequeñas
#  (p. ej. 3 etapas y 5 trabajos) que se quieran validar en LEKIN.
#
#  USO EN GOOGLE COLAB
#      1) Ejecute la celda que contiene este código.
#      2) ruta = subir_base_datos()          # botón para elegir el Excel
#      3) resultados = ejecutar(ruta)        # reglas de despacho + AG
#
#  USO LOCAL (terminal)
#      python flow_shop_ga.py --datos datos/base_datos_flow_shop.xlsx
# =============================================================================

from __future__ import annotations

import argparse
import os
import time
import warnings
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
import matplotlib.pyplot as plt


# =============================================================================
# 1. PARÁMETROS POR DEFECTO
#    Cualquiera de ellos se puede cambiar en la hoja "Parametros" de la base de
#    datos o al llamar ejecutar(ruta, generaciones=500, ...).
# =============================================================================

PARAMETROS_DEFECTO = {
    # --- unidades de los datos
    "unidad_tiempo": "min/docena",   # unidad de la hoja Tiempos: min/docena, min/par, h/docena o h/par
    "pares_por_docena": 12,
    # --- lot streaming
    "max_sublotes": 4,               # máximo de sublotes por trabajo (si no se da en la hoja Trabajos)
    "tam_min_sublote": 12,           # tamaño mínimo de un sublote [pares]
    # --- algoritmo genético
    "tam_poblacion": 80,
    "generaciones": 300,
    "prob_cruce": 0.9,
    "prob_mutacion": 0.3,
    "elitismo": 2,
    "tam_torneo": 3,
    "max_sin_mejora": 100,           # generaciones sin mejora antes de detenerse
    "semilla": 42,
    # --- reglas de despacho
    "sublotes_reglas": 1,            # 1 = lote completo (comparable con LEKIN); >1 = sublotes iguales
}

UMBRAL_CERO = 0.15  # llave de tamaño por debajo de este valor => el sublote no se usa


# =============================================================================
# 2. LECTURA DE LA BASE DE DATOS
# =============================================================================

@dataclass
class Instancia:
    trabajos: list[str]               # identificador de cada trabajo
    demanda: np.ndarray               # (n,) pares por trabajo
    fecha_entrega: np.ndarray         # (n,) d_j [h]
    max_sublotes: np.ndarray          # (n,) sublotes máximos por trabajo
    tiempo_par: np.ndarray            # (n, K) horas por par en cada etapa (máquina de velocidad 1)
    etapas: list[str]                 # nombre de cada etapa
    tipos: list[str]                  # "P" (idénticas) o "Q" (uniformes)
    velocidades: list[np.ndarray]     # por etapa: velocidad relativa de cada máquina
    maquinas: list[list[str]]         # por etapa: nombre de cada máquina

    @property
    def n(self) -> int:
        return len(self.trabajos)

    @property
    def K(self) -> int:
        return len(self.etapas)

    def notacion(self) -> str:
        campos = ", ".join(f"({t}{len(v)})" for t, v in zip(self.tipos, self.velocidades))
        return f"FH{self.K}, ({campos}) | ls | ΣTj"

    def resumen(self) -> str:
        lineas = [self.notacion(),
                  f"Trabajos: {self.n}   Demanda total: {int(self.demanda.sum())} pares"]
        for k in range(self.K):
            carga = (self.demanda * self.tiempo_par[:, k]).sum() / self.velocidades[k].sum()
            lineas.append(f"  Etapa {k + 1} {self.etapas[k]:<12} {self.tipos[k]}{len(self.velocidades[k]):<3}"
                          f" carga media por máquina = {carga:8.2f} h")
        lineas.append(f"Fechas de entrega: {self.fecha_entrega.min():.1f} a {self.fecha_entrega.max():.1f} h")
        return "\n".join(lineas)


def _normalizar(texto) -> str:
    t = str(texto).strip().lower()
    for a, b in (("á", "a"), ("é", "e"), ("í", "i"), ("ó", "o"), ("ú", "u"), ("ñ", "n"), (" ", "_")):
        t = t.replace(a, b)
    return t


def _hoja(libro: dict, nombre: str, obligatoria: bool = True) -> pd.DataFrame | None:
    for k, df in libro.items():
        if _normalizar(k) == _normalizar(nombre):
            df = df.dropna(how="all").copy()
            df.columns = [_normalizar(c) for c in df.columns]
            return df
    if obligatoria:
        raise ValueError(f"La base de datos no tiene la hoja '{nombre}'. Hojas encontradas: {list(libro)}")
    return None


def _columnas(df: pd.DataFrame, cols: list[str], hoja: str) -> None:
    faltan = [c for c in cols if c not in df.columns]
    if faltan:
        raise ValueError(f"En la hoja '{hoja}' faltan las columnas {faltan}. Encontradas: {list(df.columns)}")


def _id(valor) -> str:
    """Identificador de trabajo como texto (1.0 -> '1')."""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def leer_parametros(libro: dict) -> dict:
    params = dict(PARAMETROS_DEFECTO)
    df = _hoja(libro, "Parametros", obligatoria=False)
    if df is None:
        return params
    _columnas(df, ["parametro", "valor"], "Parametros")
    for _, fila in df.iterrows():
        clave = _normalizar(fila["parametro"])
        if clave not in params:
            warnings.warn(f"Parámetro desconocido '{fila['parametro']}' (se ignora).")
            continue
        if pd.isna(fila["valor"]):
            continue
        tipo = type(PARAMETROS_DEFECTO[clave])
        params[clave] = str(fila["valor"]).strip() if tipo is str else tipo(float(fila["valor"]))
    return params


def factor_horas_por_par(unidad: str, pares_por_docena: int) -> float:
    unidades = {
        "min/docena": 1 / (60 * pares_por_docena),
        "min/par": 1 / 60,
        "h/docena": 1 / pares_por_docena,
        "h/par": 1.0,
    }
    u = unidad.replace(" ", "").lower()
    if u not in unidades:
        raise ValueError(f"unidad_tiempo '{unidad}' no válida. Use una de {list(unidades)}")
    return unidades[u]


def cargar_base_datos(ruta: str) -> tuple[Instancia, dict]:
    """Lee la base de datos de Excel y devuelve (instancia, parámetros)."""
    libro = pd.read_excel(ruta, sheet_name=None)
    params = leer_parametros(libro)

    # ---------- Etapas
    et = _hoja(libro, "Etapas")
    _columnas(et, ["etapa", "nombre", "tipo", "num_maquinas"], "Etapas")
    et["etapa"] = et["etapa"].astype(int)
    et = et.sort_values("etapa")
    K = len(et)
    if et["etapa"].tolist() != list(range(1, K + 1)):
        raise ValueError("Las etapas deben numerarse 1, 2, 3, ... sin saltos.")
    nombres_etapas = et["nombre"].astype(str).str.strip().tolist()
    tipos = et["tipo"].astype(str).str.strip().str.upper().tolist()
    num_maq = et["num_maquinas"].astype(int).tolist()
    if any(t not in ("P", "Q") for t in tipos):
        raise ValueError("La columna 'tipo' de la hoja Etapas solo admite P (idénticas) o Q (uniformes).")

    # ---------- Máquinas (opcional)
    ma = _hoja(libro, "Maquinas", obligatoria=False)
    velocidades, maquinas = [], []
    for k in range(K):
        sub = None
        if ma is not None and len(ma):
            _columnas(ma, ["etapa", "maquina"], "Maquinas")
            sub = ma[ma["etapa"].astype(int) == k + 1]
        if sub is None or sub.empty:
            velocidades.append(np.ones(num_maq[k]))
            maquinas.append([f"{nombres_etapas[k][:3].upper()}-{i + 1:02d}" for i in range(num_maq[k])])
            continue
        if len(sub) != num_maq[k]:
            raise ValueError(f"La etapa {k + 1} declara {num_maq[k]} máquinas en 'Etapas' "
                             f"pero 'Maquinas' tiene {len(sub)}.")
        v = sub["velocidad"].fillna(1.0).astype(float).to_numpy() if "velocidad" in sub else np.ones(len(sub))
        if (v <= 0).any():
            raise ValueError(f"Velocidades no positivas en la etapa {k + 1}.")
        if tipos[k] == "P" and not np.allclose(v, v[0]):
            warnings.warn(f"La etapa {k + 1} es de tipo P pero sus velocidades difieren; se usarán tal cual.")
        velocidades.append(v)
        maquinas.append(sub["maquina"].astype(str).tolist())

    # ---------- Trabajos
    tr = _hoja(libro, "Trabajos")
    _columnas(tr, ["trabajo", "demanda_pares", "fecha_entrega_h"], "Trabajos")
    trabajos = [_id(x) for x in tr["trabajo"]]
    if len(set(trabajos)) != len(trabajos):
        raise ValueError("Hay trabajos repetidos en la hoja 'Trabajos'.")
    demanda = tr["demanda_pares"].astype(float).to_numpy()
    if (demanda <= 0).any() or not np.allclose(demanda, np.round(demanda)):
        raise ValueError("'demanda_pares' debe ser un entero positivo.")
    demanda = np.round(demanda).astype(int)
    fecha = tr["fecha_entrega_h"].astype(float).to_numpy()
    if "max_sublotes" in tr.columns:
        ms = tr["max_sublotes"].fillna(params["max_sublotes"]).astype(int).to_numpy()
    else:
        ms = np.full(len(trabajos), params["max_sublotes"], dtype=int)
    ms = np.clip(ms, 1, None)
    ms = np.minimum(ms, np.maximum(1, demanda // max(1, params["tam_min_sublote"])))
    idx = {j: i for i, j in enumerate(trabajos)}

    # ---------- Tiempos
    ti = _hoja(libro, "Tiempos")
    cols = [f"etapa_{k + 1}" for k in range(K)]
    _columnas(ti, ["trabajo"] + cols, "Tiempos")
    p = np.full((len(trabajos), K), np.nan)
    for _, fila in ti.iterrows():
        j = _id(fila["trabajo"])
        if j not in idx:
            raise ValueError(f"El trabajo '{j}' de la hoja 'Tiempos' no existe en 'Trabajos'.")
        p[idx[j]] = [float(fila[c]) for c in cols]
    if np.isnan(p).any():
        faltan = [trabajos[i] for i in np.where(np.isnan(p).any(axis=1))[0]]
        raise ValueError(f"Faltan tiempos de proceso para los trabajos {faltan}.")
    if (p < 0).any():
        raise ValueError("Los tiempos de proceso no pueden ser negativos.")
    p = p * factor_horas_por_par(params["unidad_tiempo"], params["pares_por_docena"])

    inst = Instancia(trabajos, demanda, fecha, ms, p, nombres_etapas, tipos, velocidades, maquinas)
    return inst, params


# =============================================================================
# 3. PROGRAMA DE PRODUCCIÓN Y MEDIDAS DE DESEMPEÑO
# =============================================================================

@dataclass
class Programa:
    metodo: str
    operaciones: list[tuple]   # (trabajo, sublote, pares, etapa, maquina, inicio, fin)
    completacion: np.ndarray   # (n,) C_j [h]
    tardanza: np.ndarray       # (n,) T_j [h]
    tamanos: list[np.ndarray]  # por trabajo: tamaño de sus sublotes

    @property
    def tardanza_total(self) -> float:
        return float(self.tardanza.sum())

    @property
    def makespan(self) -> float:
        return float(self.completacion.max())

    def metricas(self) -> dict:
        return {
            "metodo": self.metodo,
            "tardanza_total_h": round(self.tardanza_total, 3),
            "makespan_h": round(self.makespan, 3),
            "trabajos_tardios": int((self.tardanza > 1e-9).sum()),
            "tardanza_max_h": round(float(self.tardanza.max()), 3),
            "tardanza_media_h": round(float(self.tardanza.mean()), 3),
            "flujo_medio_h": round(float(self.completacion.mean()), 3),
            "suma_completacion_h": round(float(self.completacion.sum()), 3),
        }


def _construir_programa(metodo, inst, sublotes, fin_ultima, operaciones) -> Programa:
    C = np.zeros(inst.n)
    tam = [[] for _ in range(inst.n)]
    for s, (j, r, q) in enumerate(sublotes):
        C[j] = max(C[j], fin_ultima[s])
        tam[j].append(q)
    T = np.maximum(0.0, C - inst.fecha_entrega)
    return Programa(metodo, operaciones, C, T, [np.array(t) for t in tam])


def tamanos_sublotes(claves: np.ndarray, demanda: int, s_max: int, tam_min: int) -> np.ndarray:
    """Convierte llaves en [0,1] en tamaños enteros de sublotes que suman la demanda.

    Llaves < UMBRAL_CERO anulan el sublote; los sublotes con menos de tam_min
    pares se eliminan y su cantidad se reparte entre los demás.
    """
    w = claves[:s_max].astype(float).copy()
    w[w < UMBRAL_CERO] = 0.0
    if w.sum() <= 0:
        w[np.argmax(claves[:s_max])] = 1.0
    while True:
        prop = w / w.sum() * demanda
        base = np.floor(prop).astype(int)
        resto = demanda - base.sum()
        if resto > 0:  # método del mayor residuo
            orden = np.argsort(-(prop - base) + (w == 0) * 10)
            base[orden[:resto]] += 1
        activos = base > 0
        pequenos = activos & (base < tam_min)
        if not pequenos.any() or activos.sum() <= 1:
            return base
        cand = np.where(pequenos)[0]
        w[cand[np.argmin(base[cand])]] = 0.0


# =============================================================================
# 4. REGLAS DE DESPACHO
#    Programación no retrasada (non-delay) etapa por etapa: cada vez que una
#    máquina queda libre se elige, entre los sublotes que ya están en cola, el
#    de mayor prioridad según la regla. Si hay varias máquinas libres se usa la
#    más rápida (empate: la de menor índice). Empates entre sublotes: menor
#    número de trabajo y luego menor número de sublote.
#    Con sublotes_reglas = 1 cada trabajo se procesa como un lote completo, que
#    es el caso que se puede reproducir en LEKIN (Flexible Flow Shop).
# =============================================================================

REGLAS = {
    "FIFO": "Primero en llegar a la cola de la etapa, primero en procesarse",
    "LIFO": "Último en llegar a la cola de la etapa, primero en procesarse",
    "SPT":  "Menor tiempo de proceso en la etapa",
    "LPT":  "Mayor tiempo de proceso en la etapa",
    "EDD":  "Fecha de entrega más temprana",
    "MS":   "Mínima holgura: d_j - t - trabajo restante",
    "CR":   "Menor razón crítica: (d_j - t) / trabajo restante",
    "LWKR": "Menor trabajo restante (etapa actual y siguientes)",
    "MWKR": "Mayor trabajo restante (etapa actual y siguientes)",
}


def _prioridad(regla, t, llegada, p_etapa, restante, d):
    if regla == "FIFO":
        return llegada
    if regla == "LIFO":
        return -llegada
    if regla == "SPT":
        return p_etapa
    if regla == "LPT":
        return -p_etapa
    if regla == "EDD":
        return d
    if regla == "MS":
        return d - t - restante
    if regla == "CR":
        return (d - t) / max(restante, 1e-9)
    if regla == "LWKR":
        return restante
    if regla == "MWKR":
        return -restante
    raise ValueError(f"Regla desconocida '{regla}'. Disponibles: {list(REGLAS)}")


def programar_con_regla(inst: Instancia, regla: str, sublotes_por_trabajo: int = 1,
                        tam_min: int = 1) -> Programa:
    regla = regla.upper()
    sublotes = []
    for j in range(inst.n):
        s = int(min(sublotes_por_trabajo, inst.demanda[j]))
        for r, q in enumerate(tamanos_sublotes(np.ones(s), int(inst.demanda[j]), s, tam_min)):
            if q > 0:
                sublotes.append((j, r, int(q)))
    S = len(sublotes)
    llegada = np.zeros(S)
    operaciones = []

    for k in range(inst.K):
        vel = inst.velocidades[k]
        libre = np.zeros(len(vel))
        fin = np.zeros(S)
        pendientes = set(range(S))
        while pendientes:
            t = max(libre.min(), min(llegada[s] for s in pendientes))
            cola = [s for s in pendientes if llegada[s] <= t + 1e-9]

            def clave(s):
                j, r, q = sublotes[s]
                p_et = q * inst.tiempo_par[j, k]
                rest = q * inst.tiempo_par[j, k:].sum()
                return (_prioridad(regla, t, llegada[s], p_et, rest, inst.fecha_entrega[j]), j, r)

            s = min(cola, key=clave)
            j, r, q = sublotes[s]
            libres = np.where(libre <= t + 1e-9)[0]
            m = int(libres[np.argmax(vel[libres])])  # la más rápida entre las libres
            fin[s] = t + q * inst.tiempo_par[j, k] / vel[m]
            libre[m] = fin[s]
            operaciones.append((j, r, q, k, m, float(t), float(fin[s])))
            pendientes.remove(s)
        llegada = fin  # sin tiempos de transporte

    nombre = regla if sublotes_por_trabajo == 1 else f"{regla} ({sublotes_por_trabajo} sublotes)"
    return _construir_programa(nombre, inst, sublotes, fin, operaciones)


def comparar_reglas(inst: Instancia, reglas: list[str] | None = None, sublotes_por_trabajo: int = 1,
                    tam_min: int = 1) -> tuple[pd.DataFrame, dict[str, Programa]]:
    """Ejecuta todas las reglas y devuelve una tabla comparativa y los programas."""
    programas = {}
    for regla in reglas or list(REGLAS):
        programas[regla] = programar_con_regla(inst, regla, sublotes_por_trabajo, tam_min)
    tabla = pd.DataFrame([p.metricas() for p in programas.values()])
    tabla.insert(1, "descripcion", [REGLAS[r.upper()] for r in programas])
    return tabla, programas


# =============================================================================
# 5. ALGORITMO GENÉTICO
#
#  Cromosoma de dos partes:
#   * secuencia: cada trabajo j aparece max_sublotes_j veces; la r-ésima
#     aparición es su sublote r. Define la prioridad en la etapa 1.
#   * claves: una llave en [0,1] por sublote que define su tamaño (sublotes
#     consistentes). Llaves < 0.15 anulan el sublote, así el AG también decide
#     cuántos sublotes usar.
#  Decodificación: etapa 1 en el orden de la secuencia; etapas siguientes en
#  orden de llegada (FIFO). Cada sublote va a la máquina que lo termina antes
#  (ECT), considerando la velocidad de las máquinas uniformes.
#  Operadores: torneo, cruce POX (secuencia) + uniforme por trabajo (claves),
#  mutación swap/inserción + gaussiana, elitismo, parada por estancamiento.
# =============================================================================

@dataclass
class Cromosoma:
    secuencia: np.ndarray
    claves: np.ndarray

    def copia(self) -> "Cromosoma":
        return Cromosoma(self.secuencia.copy(), self.claves.copy())


def decodificar(cr: Cromosoma, inst: Instancia, tam_min: int, detalle: bool = False):
    n, K, p = inst.n, inst.K, inst.tiempo_par
    tam = np.zeros(cr.claves.shape, dtype=int)
    for j in range(n):
        s = int(inst.max_sublotes[j])
        tam[j, :s] = tamanos_sublotes(cr.claves[j], int(inst.demanda[j]), s, tam_min)

    contador = np.zeros(n, dtype=int)
    sublotes = []
    for j in cr.secuencia:
        r = contador[j]
        contador[j] += 1
        if tam[j, r] > 0:
            sublotes.append((int(j), int(r), int(tam[j, r])))
    S = len(sublotes)
    prioridad = np.arange(S)
    listo = np.zeros(S)
    fin = np.zeros(S)
    operaciones = []
    for k in range(K):
        libre = np.zeros(len(inst.velocidades[k]))
        vel = inst.velocidades[k]
        orden = prioridad if k == 0 else np.lexsort((prioridad, listo))
        for s in orden:
            j, r, q = sublotes[s]
            inicio = np.maximum(libre, listo[s])
            termino = inicio + q * p[j, k] / vel
            m = int(np.argmin(termino))
            libre[m] = fin[s] = termino[m]
            if detalle:
                operaciones.append((j, r, q, k, m, float(inicio[m]), float(termino[m])))
        listo = fin.copy()

    if detalle:
        return _construir_programa("Algoritmo genético", inst, sublotes, fin, operaciones)
    C = np.zeros(n)
    for s, (j, _, _) in enumerate(sublotes):
        C[j] = max(C[j], fin[s])
    return float(np.maximum(0.0, C - inst.fecha_entrega).sum()), float(C.max())


@dataclass
class ResultadoAG:
    programa: Programa
    mejor: Cromosoma
    historial_mejor: list[float] = field(default_factory=list)
    historial_promedio: list[float] = field(default_factory=list)
    generaciones: int = 0
    tiempo_s: float = 0.0


class AlgoritmoGenetico:
    def __init__(self, inst: Instancia, params: dict, verbose: bool = True):
        self.inst, self.p, self.verbose = inst, params, verbose
        semilla = params.get("semilla")
        self.rng = np.random.default_rng(None if semilla is None or semilla < 0 else semilla)
        self.s_max = int(inst.max_sublotes.max())
        self.base = np.repeat(np.arange(inst.n), inst.max_sublotes)

    def _evaluar(self, cr):
        return decodificar(cr, self.inst, self.p["tam_min_sublote"])

    @staticmethod
    def _orden(fit):
        return (round(fit[0], 6), fit[1])  # ΣTj y, como desempate, Cmax

    def _aleatorio(self):
        return Cromosoma(self.rng.permutation(self.base), self.rng.random((self.inst.n, self.s_max)))

    def _heuristico(self, criterio, n_sub):
        orden = np.argsort(criterio, kind="stable")
        sec = np.concatenate([np.full(self.inst.max_sublotes[j], j) for j in orden])
        claves = np.zeros((self.inst.n, self.s_max))
        for j in range(self.inst.n):
            claves[j, :min(n_sub, self.inst.max_sublotes[j])] = 1.0
        return Cromosoma(sec, claves)

    def _torneo(self, fits):
        cand = self.rng.choice(len(fits), size=min(self.p["tam_torneo"], len(fits)), replace=False)
        return int(min(cand, key=lambda i: self._orden(fits[i])))

    @staticmethod
    def _pox(p1, p2, conjunto):
        hijo = p1.copy()
        hijo[~conjunto[p1]] = p2[~conjunto[p2]]
        return hijo

    def _cruce(self, a, b):
        conjunto = self.rng.random(self.inst.n) < 0.5
        mascara = (self.rng.random(self.inst.n) < 0.5)[:, None]
        return (Cromosoma(self._pox(a.secuencia, b.secuencia, conjunto), np.where(mascara, a.claves, b.claves)),
                Cromosoma(self._pox(b.secuencia, a.secuencia, conjunto), np.where(mascara, b.claves, a.claves)))

    def _mutar(self, cr):
        pm, sec, L = self.p["prob_mutacion"], cr.secuencia, len(cr.secuencia)
        if L > 1 and self.rng.random() < pm:
            i, j = self.rng.choice(L, size=2, replace=False)
            if self.rng.random() < 0.5:
                sec[i], sec[j] = sec[j], sec[i]
            else:
                sec[:] = np.insert(np.delete(sec, i), j, sec[i])
        if self.rng.random() < pm:
            j = self.rng.integers(self.inst.n)
            s = self.inst.max_sublotes[j]
            cr.claves[j, :s] = np.clip(cr.claves[j, :s] + self.rng.normal(0, 0.2, s), 0, 1)
        if self.rng.random() < pm * 0.3:
            j = self.rng.integers(self.inst.n)
            r = self.rng.integers(self.inst.max_sublotes[j])
            cr.claves[j, r] = 0.0 if cr.claves[j, r] >= UMBRAL_CERO else self.rng.uniform(0.5, 1)

    def ejecutar(self) -> ResultadoAG:
        t0, P, inst = time.time(), int(self.p["tam_poblacion"]), self.inst
        trabajo_total = (inst.demanda[:, None] * inst.tiempo_par).sum(axis=1)
        pob = [self._heuristico(inst.fecha_entrega, 1), self._heuristico(inst.fecha_entrega, self.s_max),
               self._heuristico(inst.fecha_entrega - trabajo_total, self.s_max)]
        pob = pob[:P] + [self._aleatorio() for _ in range(P - min(P, len(pob)))]
        fits = [self._evaluar(c) for c in pob]
        hist_m, hist_p = [], []
        mejor_val, sin_mejora, g = None, 0, 0
        for g in range(1, int(self.p["generaciones"]) + 1):
            orden = sorted(range(P), key=lambda i: self._orden(fits[i]))
            e = int(self.p["elitismo"])
            nueva, nueva_fit = [pob[i].copia() for i in orden[:e]], [fits[i] for i in orden[:e]]
            while len(nueva) < P:
                a, b = pob[self._torneo(fits)], pob[self._torneo(fits)]
                hijos = self._cruce(a, b) if self.rng.random() < self.p["prob_cruce"] else (a.copia(), b.copia())
                for h in hijos:
                    if len(nueva) < P:
                        self._mutar(h)
                        nueva.append(h)
                        nueva_fit.append(self._evaluar(h))
            pob, fits = nueva, nueva_fit
            actual = min(fits, key=self._orden)
            hist_m.append(actual[0])
            hist_p.append(float(np.mean([f[0] for f in fits])))
            if mejor_val is None or self._orden(actual) < self._orden(mejor_val):
                mejor_val, sin_mejora = actual, 0
            else:
                sin_mejora += 1
            if self.verbose and (g == 1 or g % 10 == 0):
                print(f"Gen {g:4d} | ΣTj mejor = {actual[0]:10.3f} h | Cmax = {actual[1]:9.3f} h | "
                      f"ΣTj promedio = {hist_p[-1]:10.3f} h")
            if sin_mejora >= int(self.p["max_sin_mejora"]):
                if self.verbose:
                    print(f"Parada: {self.p['max_sin_mejora']} generaciones sin mejora.")
                break
        mejor = pob[min(range(P), key=lambda i: self._orden(fits[i]))]
        prog = decodificar(mejor, inst, self.p["tam_min_sublote"], detalle=True)
        return ResultadoAG(prog, mejor, hist_m, hist_p, g, time.time() - t0)


# =============================================================================
# 6. TABLAS, GRÁFICOS Y EXPORTACIÓN
# =============================================================================

def tabla_trabajos(prog: Programa, inst: Instancia) -> pd.DataFrame:
    return pd.DataFrame({
        "trabajo": inst.trabajos,
        "demanda_pares": inst.demanda,
        "num_sublotes": [len(t) for t in prog.tamanos],
        "tamanos_sublotes": [", ".join(map(str, t)) for t in prog.tamanos],
        "fecha_entrega_h": inst.fecha_entrega,
        "completacion_h": np.round(prog.completacion, 3),
        "tardanza_h": np.round(prog.tardanza, 3),
    })


def tabla_operaciones(prog: Programa, inst: Instancia) -> pd.DataFrame:
    df = pd.DataFrame([{
        "trabajo": inst.trabajos[j], "sublote": r + 1, "pares": q,
        "etapa": k + 1, "nombre_etapa": inst.etapas[k], "maquina": inst.maquinas[k][m],
        "inicio_h": round(a, 3), "fin_h": round(b, 3), "duracion_h": round(b - a, 3),
    } for j, r, q, k, m, a, b in prog.operaciones])
    return df.sort_values(["etapa", "maquina", "inicio_h"], ignore_index=True)


def grafico_gantt(prog: Programa, inst: Instancia, ruta: str | None = None, mostrar: bool = True):
    ops = tabla_operaciones(prog, inst)
    cmap = plt.get_cmap("tab20")
    color = {t: cmap(i % 20) for i, t in enumerate(inst.trabajos)}
    filas = [(k + 1, m) for k in range(inst.K) for m in inst.maquinas[k]]
    y = {f: i for i, f in enumerate(filas)}
    fig, ax = plt.subplots(figsize=(16, max(5, 0.3 * len(filas))))
    for _, o in ops.iterrows():
        yi = y[(o["etapa"], o["maquina"])]
        ax.barh(yi, o["duracion_h"], left=o["inicio_h"], height=0.8, color=color[o["trabajo"]],
                edgecolor="black", linewidth=0.3)
        if o["duracion_h"] > prog.makespan * 0.03:
            ax.text(o["inicio_h"] + o["duracion_h"] / 2, yi, f"{o['trabajo']}.{o['sublote']}",
                    ha="center", va="center", fontsize=6)
    ax.set_yticks(range(len(filas)))
    ax.set_yticklabels([f"E{e}-{m}" for e, m in filas], fontsize=6)
    ax.invert_yaxis()
    ax.set_xlabel("Tiempo (h)")
    ax.set_title(f"{prog.metodo} — ΣTj = {prog.tardanza_total:.2f} h, Cmax = {prog.makespan:.2f} h")
    ax.grid(axis="x", alpha=0.3)
    ax.legend([plt.Rectangle((0, 0), 1, 1, color=color[t]) for t in inst.trabajos], inst.trabajos,
              loc="upper left", bbox_to_anchor=(1.01, 1), fontsize=7, title="Trabajo")
    fig.tight_layout()
    if ruta:
        fig.savefig(ruta, dpi=150)
    if mostrar:
        plt.show()
    plt.close(fig)


def grafico_convergencia(res: ResultadoAG, ruta: str | None = None, mostrar: bool = True):
    fig, ax = plt.subplots(figsize=(8, 4.5))
    g = np.arange(1, len(res.historial_mejor) + 1)
    ax.plot(g, res.historial_mejor, lw=2, label="Mejor individuo")
    ax.plot(g, res.historial_promedio, lw=1, alpha=0.7, label="Promedio de la población")
    ax.set_xlabel("Generación")
    ax.set_ylabel("Tardanza total ΣTj (h)")
    ax.set_title("Convergencia del algoritmo genético")
    ax.grid(alpha=0.3)
    ax.legend()
    fig.tight_layout()
    if ruta:
        fig.savefig(ruta, dpi=150)
    if mostrar:
        plt.show()
    plt.close(fig)


def exportar_excel(ruta: str, inst: Instancia, params: dict, comparacion: pd.DataFrame,
                   programas: dict[str, Programa], res_ag: ResultadoAG | None) -> None:
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        comparacion.to_excel(w, sheet_name="Comparacion", index=False)
        # completación y tardanza de cada trabajo con cada método (útil para validar con LEKIN)
        cj = pd.DataFrame({"trabajo": inst.trabajos, "fecha_entrega_h": inst.fecha_entrega})
        for nombre, prog in programas.items():
            cj[f"C_{nombre}"] = np.round(prog.completacion, 3)
            cj[f"T_{nombre}"] = np.round(prog.tardanza, 3)
        cj.to_excel(w, sheet_name="Cj_Tj_por_metodo", index=False)
        # tiempo de proceso del lote completo en cada etapa [min], para cargar la instancia en LEKIN
        lekin = pd.DataFrame({"trabajo": inst.trabajos, "fecha_entrega_h": inst.fecha_entrega,
                              "fecha_entrega_min": np.round(inst.fecha_entrega * 60, 2)})
        for k in range(inst.K):
            lekin[f"etapa_{k + 1}_{inst.etapas[k]}_min"] = np.round(inst.demanda * inst.tiempo_par[:, k] * 60, 2)
        lekin.to_excel(w, sheet_name="Tiempos_lote_LEKIN", index=False)
        for nombre, prog in programas.items():
            hoja = ("AG" if nombre == "AG" else f"Regla_{nombre}")[:31]
            tabla_operaciones(prog, inst).to_excel(w, sheet_name=hoja, index=False)
        if res_ag is not None:
            tabla_trabajos(res_ag.programa, inst).to_excel(w, sheet_name="AG_Trabajos", index=False)
            pd.DataFrame({"generacion": np.arange(1, len(res_ag.historial_mejor) + 1),
                          "mejor_tardanza_total_h": res_ag.historial_mejor,
                          "promedio_tardanza_total_h": res_ag.historial_promedio}
                         ).to_excel(w, sheet_name="AG_Convergencia", index=False)
        pd.DataFrame({"parametro": list(params), "valor": [str(v) for v in params.values()]}
                     ).to_excel(w, sheet_name="Parametros_usados", index=False)


# =============================================================================
# 7. EJECUCIÓN COMPLETA
# =============================================================================

def _en_notebook() -> bool:
    try:
        return get_ipython().__class__.__name__ in ("ZMQInteractiveShell", "Shell")  # noqa: F821
    except NameError:
        return False


def ejecutar(ruta: str, carpeta: str = "resultados", ejecutar_reglas: bool = True, ejecutar_ag: bool = True,
             reglas: list[str] | None = None, mostrar_graficos: bool | None = None, **cambios) -> dict:
    """Lee la base de datos, aplica las reglas de despacho y el AG, y guarda los resultados.

    Cualquier parámetro de PARAMETROS_DEFECTO se puede cambiar aquí, p. ej.
    ejecutar(ruta, generaciones=500, tam_poblacion=100, sublotes_reglas=3).
    """
    inst, params = cargar_base_datos(ruta)
    for k, v in cambios.items():
        if k not in params:
            raise ValueError(f"Parámetro desconocido '{k}'. Disponibles: {list(params)}")
        params[k] = v
    mostrar = _en_notebook() if mostrar_graficos is None else mostrar_graficos
    os.makedirs(carpeta, exist_ok=True)

    print("=" * 80)
    print(inst.resumen())
    print("=" * 80)

    comparacion, programas = pd.DataFrame(), {}
    if ejecutar_reglas:
        comparacion, programas = comparar_reglas(inst, reglas, params["sublotes_reglas"], params["tam_min_sublote"])
        print("\nREGLAS DE DESPACHO")
        print(comparacion.drop(columns="descripcion").to_string(index=False))

    res_ag = None
    if ejecutar_ag:
        print("\nALGORITMO GENÉTICO")
        res_ag = AlgoritmoGenetico(inst, params).ejecutar()
        programas["AG"] = res_ag.programa
        comparacion = pd.concat([comparacion, pd.DataFrame([{
            **res_ag.programa.metricas(), "descripcion": "Algoritmo genético con lot streaming"}])],
            ignore_index=True)
        print(f"ΣTj = {res_ag.programa.tardanza_total:.3f} h | Cmax = {res_ag.programa.makespan:.3f} h | "
              f"{res_ag.generaciones} generaciones en {res_ag.tiempo_s:.1f} s")
        print(tabla_trabajos(res_ag.programa, inst).to_string(index=False))
        grafico_convergencia(res_ag, os.path.join(carpeta, "convergencia_AG.png"), mostrar)
        grafico_gantt(res_ag.programa, inst, os.path.join(carpeta, "gantt_AG.png"), mostrar)

    if not comparacion.empty:
        comparacion = comparacion.sort_values(["tardanza_total_h", "makespan_h"], ignore_index=True)
        print("\nCOMPARACIÓN (ordenada por tardanza total)")
        print(comparacion[["metodo", "tardanza_total_h", "makespan_h", "trabajos_tardios",
                           "tardanza_max_h", "flujo_medio_h"]].to_string(index=False))
        if ejecutar_reglas:
            mejor_regla = next(m for m in comparacion["metodo"] if m != "AG" and m != "Algoritmo genético")
            clave = mejor_regla.split(" ")[0]
            grafico_gantt(programas[clave], inst, os.path.join(carpeta, f"gantt_{clave}.png"), mostrar)

    ruta_excel = os.path.join(carpeta, "resultados.xlsx")
    exportar_excel(ruta_excel, inst, params, comparacion, programas, res_ag)
    print(f"\nResultados guardados en la carpeta '{carpeta}' (resultados.xlsx y gráficos .png)")
    return {"instancia": inst, "parametros": params, "comparacion": comparacion,
            "programas": programas, "ag": res_ag, "excel": ruta_excel}


# =============================================================================
# 8. AYUDAS PARA GOOGLE COLAB
# =============================================================================

def subir_base_datos() -> str:
    """Muestra el botón 'Elegir archivos' de Colab y devuelve la ruta del Excel subido."""
    from google.colab import files  # solo existe en Colab
    subidos = files.upload()
    if not subidos:
        raise ValueError("No se subió ningún archivo.")
    nombre = next(iter(subidos))
    with open(nombre, "wb") as f:
        f.write(subidos[nombre])
    print(f"Base de datos cargada: {nombre}")
    return nombre


def descargar_resultados(carpeta: str = "resultados") -> None:
    """Comprime la carpeta de resultados y la descarga (Colab)."""
    import shutil
    from google.colab import files
    zip_path = shutil.make_archive(carpeta, "zip", carpeta)
    files.download(zip_path)


# =============================================================================
# 9. CREACIÓN DE BASES DE DATOS DE EJEMPLO (datos sintéticos)
# =============================================================================

LEEME = [
    ("Etapas", "etapa", "Número de la etapa en el orden del proceso (1, 2, 3, ...). Define la ruta del flujo."),
    ("Etapas", "nombre", "Nombre de la etapa (Troquelado, Ensamble, Costura, Volteado)."),
    ("Etapas", "tipo", "P = máquinas paralelas idénticas; Q = máquinas paralelas uniformes (distinta velocidad)."),
    ("Etapas", "num_maquinas", "Cantidad de máquinas en la etapa. Debe coincidir con las filas de la hoja Maquinas."),
    ("Maquinas", "etapa", "Etapa a la que pertenece la máquina."),
    ("Maquinas", "maquina", "Nombre o código de la máquina."),
    ("Maquinas", "velocidad", "Velocidad relativa. Tiempo real = tiempo estándar / velocidad. Use 1 en etapas tipo P."),
    ("Trabajos", "trabajo", "Identificador del trabajo (número o texto). Una fila por trabajo."),
    ("Trabajos", "demanda_pares", "Demanda del trabajo en pares de guantes (entero)."),
    ("Trabajos", "fecha_entrega_h", "Fecha de entrega en horas, medida desde el inicio de la programación (t = 0)."),
    ("Trabajos", "max_sublotes", "Máximo de sublotes en que el AG puede dividir el trabajo (opcional)."),
    ("Tiempos", "trabajo", "Identificador del trabajo (igual al de la hoja Trabajos)."),
    ("Tiempos", "etapa_1 ... etapa_K", "Tiempo de proceso del trabajo en cada etapa, en la unidad de 'unidad_tiempo' "
                                        "(por defecto min/docena). En etapas Q es el tiempo de una máquina de velocidad 1."),
    ("Parametros", "parametro / valor", "Configuración del modelo y del algoritmo genético. Si falta una fila se usa el "
                                        "valor por defecto."),
]

DESCRIPCION_PARAMETROS = {
    "unidad_tiempo": "Unidad de la hoja Tiempos: min/docena, min/par, h/docena o h/par",
    "pares_por_docena": "Pares que contiene una docena",
    "max_sublotes": "Máximo de sublotes por trabajo cuando no se indica en la hoja Trabajos",
    "tam_min_sublote": "Tamaño mínimo de un sublote [pares]",
    "tam_poblacion": "AG: número de individuos de la población",
    "generaciones": "AG: número máximo de generaciones",
    "prob_cruce": "AG: probabilidad de cruce",
    "prob_mutacion": "AG: probabilidad de mutación",
    "elitismo": "AG: mejores individuos que pasan intactos a la siguiente generación",
    "tam_torneo": "AG: individuos que compiten en cada torneo de selección",
    "max_sin_mejora": "AG: generaciones sin mejora para detener la búsqueda",
    "semilla": "Semilla aleatoria (mismo valor = mismos resultados; -1 = aleatoria)",
    "sublotes_reglas": "Reglas de despacho: 1 = lote completo (como LEKIN); >1 = sublotes iguales",
}


def guardar_base_datos(ruta: str, etapas: pd.DataFrame, maquinas: pd.DataFrame, trabajos: pd.DataFrame,
                       tiempos: pd.DataFrame, params: dict | None = None) -> None:
    from openpyxl.comments import Comment
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    p = dict(PARAMETROS_DEFECTO, **(params or {}))
    df_par = pd.DataFrame({"parametro": list(p), "valor": list(p.values()),
                           "descripcion": [DESCRIPCION_PARAMETROS[k] for k in p]})
    hojas = {"LEEME": pd.DataFrame(LEEME, columns=["hoja", "columna", "significado"]),
             "Etapas": etapas, "Maquinas": maquinas, "Trabajos": trabajos, "Tiempos": tiempos,
             "Parametros": df_par}
    ayuda = {(h, c): s for h, c, s in LEEME}
    os.makedirs(os.path.dirname(os.path.abspath(ruta)), exist_ok=True)
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        for nombre, df in hojas.items():
            df.to_excel(w, sheet_name=nombre, index=False)
            ws = w.sheets[nombre]
            for celda in ws[1]:
                celda.font = Font(bold=True, color="FFFFFF")
                celda.fill = PatternFill("solid", fgColor="305496")
                texto = ayuda.get((nombre, celda.value))
                if texto is None and nombre == "Tiempos" and str(celda.value).startswith("etapa_"):
                    texto = ayuda[("Tiempos", "etapa_1 ... etapa_K")]
                if texto:
                    celda.comment = Comment(texto, "Modelo FH")
            for i, col in enumerate(ws.columns, start=1):
                ancho = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(12, ancho + 2), 100)
            ws.freeze_panes = "A2"


def crear_base_datos_ejemplo(ruta: str = "base_datos_flow_shop.xlsx", n_trabajos: int = 20,
                             semilla: int = 7) -> str:
    """Base de datos sintética FH4, ((P4), (Q20), (P9), (P8))."""
    rng = np.random.default_rng(semilla)
    etapas = pd.DataFrame({"etapa": [1, 2, 3, 4],
                           "nombre": ["Troquelado", "Ensamble", "Costura", "Volteado"],
                           "tipo": ["P", "Q", "P", "P"],
                           "num_maquinas": [4, 20, 9, 8]})
    filas = []
    for _, e in etapas.iterrows():
        for i in range(e["num_maquinas"]):
            v = round(float(rng.uniform(0.85, 1.20)), 2) if e["tipo"] == "Q" else 1.0
            filas.append({"etapa": e["etapa"], "maquina": f"{e['nombre'][:3].upper()}-{i + 1:02d}", "velocidad": v})
    maquinas = pd.DataFrame(filas)
    ids = list(range(1, n_trabajos + 1))
    trabajos = pd.DataFrame({"trabajo": ids,
                             "demanda_pares": (rng.integers(25, 126, n_trabajos) * 12),
                             "fecha_entrega_h": rng.integers(12, 45, n_trabajos) * 5,
                             "max_sublotes": 4})
    tiempos = pd.DataFrame({"trabajo": ids,
                            "etapa_1": rng.integers(8, 16, n_trabajos),     # troquelado [min/docena]
                            "etapa_2": rng.integers(80, 121, n_trabajos),   # ensamble
                            "etapa_3": rng.integers(40, 56, n_trabajos),    # costura
                            "etapa_4": rng.integers(30, 46, n_trabajos)})   # volteado
    guardar_base_datos(ruta, etapas, maquinas, trabajos, tiempos)
    return ruta


def crear_instancia_pequena(ruta: str = "instancia_pequena_5x3.xlsx") -> str:
    """Instancia de 5 trabajos y 3 etapas (datos del escenario anterior) para validar en LEKIN."""
    etapas = pd.DataFrame({"etapa": [1, 2, 3], "nombre": ["Troquelado", "Ensamble", "Costura"],
                           "tipo": ["P", "P", "P"], "num_maquinas": [4, 10, 4]})
    maquinas = pd.DataFrame([{"etapa": e, "maquina": f"{n[:3].upper()}-{i + 1:02d}", "velocidad": 1.0}
                             for e, n, m in zip(etapas["etapa"], etapas["nombre"], etapas["num_maquinas"])
                             for i in range(m)])
    trabajos = pd.DataFrame({"trabajo": [1, 2, 3, 4, 5],
                             "demanda_pares": [1272, 628, 629, 660, 423],
                             "fecha_entrega_h": [200, 150, 120, 160, 180],
                             "max_sublotes": 4})
    tiempos = pd.DataFrame({"trabajo": [1, 2, 3, 4, 5],
                            "etapa_1": [30, 25, 33, 28, 36],
                            "etapa_2": [50, 45, 40, 44, 55],
                            "etapa_3": [50, 50, 46, 52, 52]})
    guardar_base_datos(ruta, etapas, maquinas, trabajos, tiempos)
    return ruta


# =============================================================================
# 10. LÍNEA DE COMANDOS (uso local)
# =============================================================================

def _cli() -> None:
    ap = argparse.ArgumentParser(description="Flow shop híbrido con algoritmo genético y reglas de despacho")
    ap.add_argument("--datos", default="datos/base_datos_flow_shop.xlsx", help="Base de datos en Excel")
    ap.add_argument("--salida", default="resultados", help="Carpeta de resultados")
    ap.add_argument("--solo-reglas", action="store_true", help="Ejecutar solo las reglas de despacho")
    ap.add_argument("--solo-ag", action="store_true", help="Ejecutar solo el algoritmo genético")
    ap.add_argument("--crear-ejemplos", action="store_true",
                    help="Crear las bases de datos de ejemplo en la carpeta datos/ y salir")
    a, _ = ap.parse_known_args()
    if a.crear_ejemplos:
        print(crear_base_datos_ejemplo("datos/base_datos_flow_shop.xlsx"))
        print(crear_instancia_pequena("datos/instancia_pequena_5x3.xlsx"))
        return
    ejecutar(a.datos, a.salida, ejecutar_reglas=not a.solo_ag, ejecutar_ag=not a.solo_reglas,
             mostrar_graficos=False)


if __name__ == "__main__" and not _en_notebook():
    _cli()
